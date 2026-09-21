"""Spreadsheets: budgets, roadmaps, headcount plans.

Rows are banded into groups with the header row repeated on each band. A lone
row ("Gripper redesign | Q3 | 240000") is unretrievable noise; the same row under
its header is answerable content. The anchor is a real cell range, so a citation
resolves to a rectangle a reader can select in Excel.
"""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.ingestion.common import Document, Section

ROWS_PER_BAND = 12


def _row_values(row) -> list[str]:
    out = []
    for cell in row:
        value = cell.value
        if value is None:
            out.append("")
        elif isinstance(value, float) and value.is_integer():
            out.append(str(int(value)))
        else:
            out.append(str(value).strip())
    while out and not out[-1]:
        out.pop()
    return out


def parse_xlsx(path: Path, source_path: str | None = None) -> Document:
    workbook = load_workbook(str(path), data_only=True, read_only=False)

    sections: list[Section] = []

    for sheet in workbook.worksheets:
        rows: list[tuple[int, list[str]]] = []
        for row in sheet.iter_rows():
            if not row:
                continue
            values = _row_values(row)
            if any(values):
                rows.append((row[0].row, values))
        if not rows:
            continue

        header_rownum, header = rows[0]
        header_line = " | ".join(header)
        width = max(len(values) for _, values in rows)
        last_col = get_column_letter(max(width, 1))

        body = rows[1:]
        if not body:
            # Header-only sheet: the header IS the content.
            sections.append(
                Section(
                    text=header_line,
                    kind="cells",
                    locator_start=f"A{header_rownum}",
                    locator_end=f"{last_col}{header_rownum}",
                    scope=sheet.title,
                )
            )
            continue

        # The header is repeated as context on every band but is not part of the
        # band's cell range, so it is named in the scope instead of the locator --
        # a citation that claimed rows 14-25 contained the header would be wrong.
        scope = f"{sheet.title} (header row {header_rownum})"
        for start in range(0, len(body), ROWS_PER_BAND):
            band = body[start : start + ROWS_PER_BAND]
            lines = [header_line] + [" | ".join(values) for _, values in band]
            sections.append(
                Section(
                    text="\n".join(lines),
                    kind="cells",
                    locator_start=f"A{band[0][0]}",
                    locator_end=f"{last_col}{band[-1][0]}",
                    scope=scope,
                )
            )

    props = workbook.properties
    authors = [props.creator.strip()] if props.creator and props.creator.strip() else []
    date = None
    stamp = props.modified or props.created
    if stamp is not None:
        date = stamp.date().isoformat()
    title = (props.title or "").strip() or path.stem

    workbook.close()

    raw_text = "\n".join(s.text for s in sections)
    return Document(
        source_path=source_path or str(path),
        source_type="xlsx",
        raw_text=raw_text,
        sections=sections,
        title=title,
        attendees_or_author=authors,
        date=date,
    )
