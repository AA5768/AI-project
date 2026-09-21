"""PowerPoint decks: status decks, roadmap reviews, QBR slides.

One section per slide body plus a separate section for speaker notes -- notes are
often where the caveat lives ("numbers are pre-audit"), and separating them keeps
that caveat citable on its own.
"""

import re
from pathlib import Path

from pptx import Presentation

from app.ingestion.common import Document, Section

AUTHOR_LINE_RE = re.compile(r"^\s*(?:author|presented by|prepared by|owner)\s*:\s*(.+)$", re.I)


def _split_people(value: str) -> list[str]:
    parts = re.split(r",| and |;|/", value)
    return [p.strip() for p in parts if p.strip()]


def _slide_title(slide) -> str | None:
    try:
        if slide.shapes.title is not None and slide.shapes.title.text.strip():
            return slide.shapes.title.text.strip()
    except (AttributeError, ValueError):
        pass
    return None


def parse_pptx(path: Path, source_path: str | None = None) -> Document:
    prs = Presentation(str(path))

    sections: list[Section] = []
    authors: list[str] = []
    title: str | None = None

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_title = _slide_title(slide)
        if slide_num == 1 and slide_title:
            title = slide_title

        body: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text and text != slide_title:
                    body.append(text)
                    matched = AUTHOR_LINE_RE.search(text)
                    if matched:
                        authors.extend(_split_people(matched.group(1)))
            if getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        rows.append(" | ".join(cells))
                if rows:
                    body.append("\n".join(rows))

        header = f"{slide_title}\n" if slide_title else ""
        if body or slide_title:
            sections.append(
                Section(
                    text=(header + "\n".join(body)).strip(),
                    kind="slide",
                    locator_start=str(slide_num),
                    locator_end=str(slide_num),
                    scope=None,
                )
            )

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                sections.append(
                    Section(
                        text=notes,
                        kind="notes",
                        locator_start=str(slide_num),
                        locator_end=str(slide_num),
                        scope=None,
                    )
                )

    props = prs.core_properties
    if not authors and props.author:
        authors = _split_people(props.author)
    date = None
    stamp = props.modified or props.created
    if stamp is not None:
        date = stamp.date().isoformat()
    if title is None:
        title = (props.title or "").strip() or path.stem

    raw_text = "\n\n".join(s.text for s in sections)
    return Document(
        source_path=source_path or str(path),
        source_type="pptx",
        raw_text=raw_text,
        sections=sections,
        title=title,
        attendees_or_author=list(dict.fromkeys(authors)),
        date=date,
    )
