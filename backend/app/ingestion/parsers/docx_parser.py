"""Word documents: design docs, meeting minutes, status write-ups.

Paragraphs carry their nearest heading as scope so a citation reads
"Rollout plan paragraph 24" rather than a bare index. Tables are emitted
separately -- a budget table inside a design doc is exactly the kind of content
that later contradicts a spreadsheet, and it needs its own anchor to be cited.
"""

import re
from pathlib import Path

from docx import Document as DocxDocument

from app.ingestion.common import Document, Section

AUTHOR_LINE_RE = re.compile(r"^\s*(?:author|prepared by|owner|written by)\s*:\s*(.+)$", re.I)
DATE_LINE_RE = re.compile(r"^\s*(?:date|last updated|updated|revised)\s*:\s*(.+)$", re.I)


def _split_people(value: str) -> list[str]:
    parts = re.split(r",| and |;|/", value)
    return [p.strip() for p in parts if p.strip()]


def parse_docx(path: Path, source_path: str | None = None) -> Document:
    docx = DocxDocument(str(path))

    sections: list[Section] = []
    heading: str | None = None
    title: str | None = None
    authors: list[str] = []
    date: str | None = None

    for index, paragraph in enumerate(docx.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue

        style = (paragraph.style.name or "") if paragraph.style else ""
        if style.startswith("Heading") or style == "Title":
            heading = text
            if title is None:
                title = text

        matched_author = AUTHOR_LINE_RE.match(text)
        if matched_author:
            authors.extend(_split_people(matched_author.group(1)))
        matched_date = DATE_LINE_RE.match(text)
        if matched_date and date is None:
            date = matched_date.group(1).strip()

        sections.append(
            Section(
                text=text,
                kind="paragraph",
                locator_start=str(index),
                locator_end=str(index),
                scope=heading,
            )
        )

    for table_index, table in enumerate(docx.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            sections.append(
                Section(
                    text="\n".join(rows),
                    kind="table",
                    locator_start=str(table_index),
                    locator_end=str(table_index),
                    scope=heading,
                )
            )

    props = docx.core_properties
    if not authors and props.author:
        authors = _split_people(props.author)
    if date is None:
        stamp = props.modified or props.created
        if stamp is not None:
            date = stamp.date().isoformat()
    if title is None:
        title = (props.title or "").strip() or path.stem

    raw_text = "\n".join(s.text for s in sections)
    return Document(
        source_path=source_path or str(path),
        source_type="docx",
        raw_text=raw_text,
        sections=sections,
        title=title,
        attendees_or_author=list(dict.fromkeys(authors)),
        date=date,
    )
