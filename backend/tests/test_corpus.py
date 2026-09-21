"""The corpus checked against its own arithmetic and its own cast.

These tests read data/ directly -- no database, no embedding model, no API key --
because the defects they catch are the kind that produce a confidently wrong
answer rather than an error. A spreadsheet whose total does not equal its line
items still parses, still chunks, still embeds, and still gets cited; the only
thing wrong with it is that it is wrong. Likewise a person who exists in a
document but not in data/people.yaml still routes -- to nobody real.

Anything deliberately inconsistent is asserted here *as* deliberate, so that a
later edit which "fixes" it fails the suite instead of quietly removing the
contradiction Part 2 exists to surface.
"""

import re

import frontmatter
import pytest
import yaml
from openpyxl import load_workbook

from app.config import PROJECT_ROOT
from app.ingestion.parsers.transcript_parser import NOT_SPEAKERS, SPEAKER_RE

DATA = PROJECT_ROOT / "data"
OFFICE = DATA / "office"
TRANSCRIPTS = DATA / "transcripts"


def _sheet(path, name):
    workbook = load_workbook(str(path), data_only=True)
    try:
        return [list(row) for row in workbook[name].iter_rows(values_only=True)]
    finally:
        workbook.close()


def _column(rows, header, label_column=0):
    """{row label: value} for one named column of a header-plus-body sheet."""
    index = rows[0].index(header)
    return {row[label_column]: row[index] for row in rows[1:] if row[label_column]}


# ----------------------------------------------------------------- arithmetic


def test_fy26_rnd_total_equals_its_line_items():
    rows = _sheet(OFFICE / "fy26-budget.xlsx", "R&D Spend")
    for column in ("Budget (USD)", "Forecast (USD)"):
        values = _column(rows, column)
        total = values.pop("R&D total")
        assert sum(values.values()) == total, f"{column} line items do not sum to the total"


def test_fy26_rnd_overrun_note_matches_the_arithmetic():
    rows = _sheet(OFFICE / "fy26-budget.xlsx", "R&D Spend")
    budget = _column(rows, "Budget (USD)")["R&D total"]
    forecast = _column(rows, "Forecast (USD)")["R&D total"]
    assert forecast - budget == 505_000  # the note on that row says "505k over plan"


def test_fy26_revenue_plan_total_equals_its_quarters():
    rows = _sheet(OFFICE / "fy26-budget.xlsx", "Revenue Plan")
    plan = _column(rows, "Plan (USD)")
    total = plan.pop("FY26 total")
    assert sum(plan.values()) == total

    actual = _column(rows, "Actual (USD)")
    variance = _column(rows, "Variance")
    for quarter, planned in plan.items():
        if isinstance(actual.get(quarter), int):
            assert actual[quarter] - planned == variance[quarter], quarter


def test_efem_cost_bridge_reconciles_to_the_standard_cost_delta():
    """The four approved actions must account for the whole standard-cost move.

    This is the property that makes "how much does the COGS-down programme
    save?" answerable from the spreadsheet rather than only from the summary
    line, which is what makes a cited answer checkable.
    """
    rows = _sheet(OFFICE / "efem-cost-reduction-tracker.xlsx", "Cost Bridge")
    current, target = _column(rows, "Current (USD)"), _column(rows, "Target (USD)")

    standard_delta = current["Standard cost per EFEM"] - target["Standard cost per EFEM"]
    component_delta = sum(
        current[line] - target[line]
        for line in (
            "Harmonic drive reducers (2)",
            "Linear encoder pair",
            "Ceramic blade set",
            "Fan filter units (6)",
            "Frame weldment and machining",
            "Labour and cleanroom overhead",
        )
    )
    assert component_delta == standard_delta == 31_700


def test_efem_gross_margin_matches_list_price_minus_standard_cost():
    rows = _sheet(OFFICE / "efem-cost-reduction-tracker.xlsx", "Cost Bridge")
    current, target = _column(rows, "Current (USD)"), _column(rows, "Target (USD)")

    for column, quoted in ((current, current["Gross margin"]), (target, target["Gross margin"])):
        price = column["List price"]
        cost = column["Standard cost per EFEM"]
        computed = round((price - cost) / price * 100)
        assert f"{computed}%" == quoted


def test_efem_action_savings_match_the_cost_bridge_lines():
    tracker = OFFICE / "efem-cost-reduction-tracker.xlsx"
    actions = _column(_sheet(tracker, "Actions"), "Saving per machine (USD)")
    rows = _sheet(tracker, "Cost Bridge")
    current, target = _column(rows, "Current (USD)"), _column(rows, "Target (USD)")

    pairs = {
        "Remove three machined datums from the frame": "Frame weldment and machining",
        "Five fan filter units with redesigned plenum": "Fan filter units (6)",
        "Rev B encoder selection": "Linear encoder pair",
        "Move harness assembly out of the clean bay": "Labour and cleanroom overhead",
    }
    for action, line in pairs.items():
        assert actions[action] == current[line] - target[line], action


def test_fan_filter_cost_bridge_line_equals_the_bom_unit_cost():
    """Cross-document: the tracker's per-set figures derive from the BOM unit price.

    They were 200 apart -- 34,000 against 6 x 5,700 -- while the target on the
    same row (28,500 = 5 x 5,700) proved which one was right. Nothing errored;
    the two documents simply disagreed, and a question about fan filter cost
    could be answered correctly from either and still contradict the other.
    """
    bom = _sheet(OFFICE / "helios3-bom-long-lead.xlsx", "Long Lead Items")
    unit_cost = _column(bom, "Unit cost (USD)")["Fan filter unit, ISO Class 2"]
    quantity = _column(bom, "Qty per machine")["Fan filter unit, ISO Class 2"]

    rows = _sheet(OFFICE / "efem-cost-reduction-tracker.xlsx", "Cost Bridge")
    assert _column(rows, "Current (USD)")["Fan filter units (6)"] == unit_cost * quantity
    assert _column(rows, "Target (USD)")["Fan filter units (6)"] == unit_cost * (quantity - 1)


def test_build_rate_achievable_is_the_binding_constraint():
    rows = _sheet(OFFICE / "helios3-bom-long-lead.xlsx", "Build Rate Constraint")
    header, body = rows[0], rows[1:]
    columns = {name: header.index(name) for name in header if name}
    for row in body:
        limits = (
            row[columns["Plan (machines)"]],
            row[columns["Reducer supply (machines)"]],
            row[columns["Cleanroom capacity (machines)"]],
        )
        assert row[columns["Achievable"]] == min(limits), row[0]


def test_spreadsheet_numbers_are_stored_as_numbers():
    """A quantity written as text sums to zero and sorts lexicographically."""
    numeric_headers = re.compile(r"\((USD|machines)\)$|^Qty|^Lead time|^Credit owed", re.I)
    for path in sorted(OFFICE.glob("*.xlsx")):
        workbook = load_workbook(str(path), data_only=True)
        try:
            for sheet in workbook.worksheets:
                rows = [list(r) for r in sheet.iter_rows(values_only=True)]
                if len(rows) < 2:
                    continue
                for index, header in enumerate(rows[0]):
                    if not (header and numeric_headers.search(str(header))):
                        continue
                    for row in rows[1:]:
                        value = row[index] if index < len(row) else None
                        if value in (None, "", "TBD"):
                            continue
                        if isinstance(value, str) and value.endswith("%"):
                            # The EFEM cost bridge carries a gross-margin row in
                            # its USD columns. Ratios in a currency column are a
                            # real thing finance packs do; a quantity stored as
                            # text is not.
                            continue
                        assert isinstance(value, (int, float)), (
                            f"{path.name}/{sheet.title} '{header}' holds {value!r} as text"
                        )
        finally:
            workbook.close()


# ----------------------------------------------------------------------- cast


@pytest.fixture(scope="module")
def cast():
    return yaml.safe_load((DATA / "people.yaml").read_text(encoding="utf-8"))


def test_every_attendee_is_in_the_people_directory(cast):
    """An attendee absent from people.yaml routes to a name with no title, no
    department and no mailbox -- addressable in the response shape, not in life."""
    unknown = []
    for path in sorted(TRANSCRIPTS.glob("*.md")):
        post = frontmatter.load(path)
        attendees = post.get("attendees") or post.get("participants") or post.get("present") or []
        if isinstance(attendees, str):
            attendees = attendees.split(",")
        for attendee in attendees:
            name = re.sub(r"\s*\(.*\)$", "", str(attendee)).strip()
            if name and name not in cast:
                unknown.append(f"{path.name}: {name}")
    assert not unknown


def test_every_named_speaker_is_a_known_person(cast):
    """The one exception is registered: the deliberately unattributed standup
    notes use "TBD:" as a speaker. It must stay unresolvable -- resolve_person
    refuses to create a person for it, so it never enters the routing graph."""
    first_names = {name.split()[0] for name in cast}
    unknown = []
    for path in sorted(TRANSCRIPTS.glob("*.md")):
        for line in frontmatter.load(path).content.splitlines():
            match = SPEAKER_RE.match(line)
            if not match:
                continue
            speaker = match.group(1).strip()
            if speaker.lower() in NOT_SPEAKERS or speaker in cast or speaker in first_names:
                continue
            unknown.append(f"{path.name}: {speaker}")
    assert unknown == ["2026-05-07-eng-standup-partial-notes.md: TBD"]


def test_transcript_date_matches_its_filename():
    for path in sorted(TRANSCRIPTS.glob("*.md")):
        if not path.name[:4].isdigit():
            continue
        assert str(frontmatter.load(path).get("date")) == path.name[:10], path.name


# ------------------------------------------------------- registered on purpose


def test_the_budget_deck_is_still_stale_against_the_meetings():
    """Part 2's contradiction surfacing needs this to stay broken.

    Helios-3 NRE was cut to 3.2M on 2026-05-21 and set to 3.4M on 2026-07-23;
    fy26-budget.xlsx still carries the original 3.6M. A well-meaning edit that
    "corrects" the spreadsheet would delete the test case.
    """
    rows = _sheet(OFFICE / "fy26-budget.xlsx", "R&D Spend")
    assert _column(rows, "Budget (USD)")["Helios-3 NRE"] == 3_600_000

    workshop = (TRANSCRIPTS / "2026-07-23-efem-cogs-down-workshop.md").read_text(encoding="utf-8")
    budget_review = (TRANSCRIPTS / "2026-05-21-fy26-budget-review.md").read_text(encoding="utf-8")
    assert "3.2" in budget_review
    assert "3.4" in workshop


def test_the_cogs_workshop_rounds_where_the_tracker_is_exact():
    """Kept, not fixed: Owen quotes round numbers in a meeting.

    He says the fan filter set is 34,000 and the four actions total 31,500; the
    tracker says 34,200 and 31,700. People round. The tracker is the system of
    record and reconciles exactly (see the cost-bridge tests above), so the
    discrepancy is a realistic one between a spoken figure and a spreadsheet
    rather than an error -- and it gives correction capture something true to
    fire on.
    """
    workshop = (TRANSCRIPTS / "2026-07-23-efem-cogs-down-workshop.md").read_text(encoding="utf-8")
    assert "34,000 for the set of six" in workshop
    assert "That is 31,500" in workshop
