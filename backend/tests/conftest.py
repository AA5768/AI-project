import sys
from pathlib import Path

import pytest

BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))


@pytest.fixture
def corpus(tmp_path: Path) -> Path:
    """A miniature data/ directory exercising all four source types plus the
    awkward cases: bold speaker labels, label lines that look like speakers,
    an unattributed near-empty file, and a multi-sheet workbook."""
    return build_corpus(tmp_path)


@pytest.fixture(scope="session")
def ingested_db(tmp_path_factory) -> Path:
    """The same corpus, ingested once for the whole session.

    Part 2's tests all read the database rather than write it, so they can share
    one ingestion; re-running it per test would pay for the embedding model
    repeatedly and test nothing extra. Heuristic enrichment keeps it offline.
    """
    from app.ingestion.pipeline import run

    root = tmp_path_factory.mktemp("shared")
    data = build_corpus(root)
    db = root / "shared.db"
    run(data_dir=data, db_path=db, use_llm=False, reset=True)
    return db


def build_corpus(tmp_path: Path) -> Path:
    data = tmp_path / "data"
    (data / "transcripts").mkdir(parents=True)
    (data / "office").mkdir(parents=True)

    (data / "people.yaml").write_text(
        "Priya Raman:\n  title: VP Engineering\n  department: Engineering\n"
        "Dana Okafor:\n  title: Staff Firmware Engineer\n  department: Engineering\n",
        encoding="utf-8",
    )

    (data / "transcripts" / "standup.md").write_text(
        "---\n"
        "title: Helios-3 Program Weekly\n"
        "date: 2026-03-04\n"
        "attendees:\n  - Priya Raman\n  - Dana Okafor\n  - Marcus Lindqvist\n"
        "---\n\n"
        "## End effector status\n\n"
        "**Dana Okafor:** The edge grip end effector is blocking the pilot. Wafer\n"
        "placement repeatability drifts after six hours of continuous handling.\n\n"
        "**Priya Raman**: That's a blocker. We decided to hold the pilot ship date.\n"
        "Dana will produce an RCA by end of Q1.\n\n"
        "Action: Marcus to notify the customer.\n",
        encoding="utf-8",
    )

    (data / "transcripts" / "sparse.md").write_text(
        "---\ntitle: Quick sync\ndate: 2026-03-06\nattendees: []\n---\n\nTBD.\n",
        encoding="utf-8",
    )

    from docx import Document as Docx

    docx = Docx()
    docx.add_heading("Northwind Escalation Note", level=0)
    docx.add_paragraph("Author: Marcus Lindqvist")
    docx.add_heading("Problem", level=1)
    docx.add_paragraph(
        "Northwind reported a SEV1 outage in the OEE telemetry dashboard. "
        "Telemetry ingestion lagged 40 minutes, breaching the SLA."
    )
    table = docx.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Workstream"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "Telemetry backpressure"
    table.cell(1, 1).text = "Dana Okafor"
    docx.core_properties.author = "Marcus Lindqvist"
    docx.save(data / "office" / "escalation.docx")

    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Q1 Roadmap Review"
    slide.placeholders[1].text = "End effector redesign: 180000\nVantage 4.2 release: on track"
    slide.notes_slide.notes_text_frame.text = "Numbers are pre-audit and may be stale."
    prs.core_properties.author = "Priya Raman"
    prs.save(data / "office" / "roadmap.pptx")

    from openpyxl import Workbook

    wb = Workbook()
    sheet = wb.active
    sheet.title = "Budget"
    sheet.append(["Line item", "Quarter", "Planned", "Actual"])
    for i in range(20):
        sheet.append([f"Line {i}", f"Q{i % 4 + 1}", 10000 + i * 1000, 9500 + i * 1000])
    head = wb.create_sheet("Headcount")
    head.append(["Department", "Open reqs"])
    head.append(["Engineering", 6])
    wb.properties.creator = "Sofia Bergstrom"
    wb.save(data / "office" / "budget.xlsx")

    return data
